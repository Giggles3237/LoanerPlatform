"""Excel workbook of a weekly plan: a summary sheet plus one sheet per advisor."""

from __future__ import annotations

import io
from datetime import date

from .planner import RULE_ACTIONS
from .settings import RULE_LABELS

COLUMNS = [
    ("Day", 7), ("#", 4), ("Done", 6), ("Customer", 24), ("Phone", 16), ("Email", 28),
    ("Vehicle", 26), ("Task", 22), ("Why", 70), ("Suggested action", 36),
    ("Assigned because", 26), ("Outcome / notes", 40),
]


def tasks_for(plan: dict, advisor: str) -> list[dict]:
    order = {d: i for i, d in enumerate(plan["work_days"])}
    items = [t for t in plan["tasks"] if t["advisor"] == advisor]
    return sorted(items, key=lambda t: (order.get(t["day"], 99), -t["score"]))


def assignment_note(task: dict) -> str:
    if task["reassigned"]:
        prev = task["previous_owner"] or "no owner"
        return f"Reassigned (was: {prev})"
    return task["owner_basis"]


def plan_workbook(plan: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="262626")
    wrap = Alignment(wrap_text=True, vertical="top")

    summary = wb.active
    summary.title = "Summary"
    week = date.fromisoformat(plan["week_start"])
    summary.append([f"{plan['dealership']} — advisor task plans, week of {week:%B %d, %Y}"])
    summary["A1"].font = Font(bold=True, size=14)
    summary.append([])
    rules = list(RULE_LABELS)
    header = ["Advisor", "Tasks", "Capacity", "Own customers", "Reassigned",
              *[RULE_LABELS[r] for r in rules]]
    summary.append(header)
    for cell in summary[3]:
        cell.font, cell.fill = head_font, head_fill
    for a in plan["advisors"]:
        summary.append([a["name"], a["tasks"], a["capacity"], a["own"], a["reassigned"],
                        *[a["by_rule"].get(r, 0) for r in rules]])
    summary.column_dimensions["A"].width = 22
    for col in "BCDEFGHIJ":
        summary.column_dimensions[col].width = 16

    for a in plan["advisors"]:
        ws = wb.create_sheet(a["name"][:31])
        ws.append([c for c, _ in COLUMNS])
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
        for i, (_, width) in enumerate(COLUMNS):
            ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = width
        ws.freeze_panes = "A2"
        for n, t in enumerate(tasks_for(plan, a["name"]), start=1):
            primary = max(t["reasons"], key=lambda r: r["score"])
            email = "DO NOT EMAIL" if t["do_not_email"] else t["email"]
            ws.append([
                t["day"], n, "", t["customer"], "\n".join(t["phones"]) or "Look up in CRM",
                email, t["vehicle"], " + ".join(r["label"] for r in t["reasons"]),
                "\n".join(f"• {r['detail']}" for r in t["reasons"]),
                RULE_ACTIONS[primary["rule"]], assignment_note(t), "",
            ])
            for cell in ws[ws.max_row]:
                cell.alignment = wrap
        ws.print_options.gridLines = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:1"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
